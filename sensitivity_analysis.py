# -*- coding: utf-8 -*-
"""阈值敏感性分析 —— 回应外审意见 2。

审稿意见：基于寄存器的交叉手重新分配使用固定的音高阈值（C5=72 / C3=48），
双手拆分使用固定阈值（16 半音）。这些阈值的敏感性如何？

方法：one-at-a-time 扫描，固定其余参数为论文默认值，对全语料
（10 首、36 页、400 小节）重跑完整分析管线，计算
  - 60 单元一致率（c = 1 - |a-h|/n，n 取人工小节数）
  - 难度分与等级的 Spearman ρ
扫描维度：
  1. hand_split_threshold ∈ {12..20}（默认 16）
  2. cross_hand high（bass 谱表 > 阈值归右手）∈ {58..79}（默认 72 = C5）
  3. cross_hand low（treble 谱表 < 阈值归左手）∈ {36..55}（默认 48 = C3）

输出:
  - output_grades/reports/敏感性分析.md + .xlsx
  - cvaa_figs_en/fig_sensitivity_en.png（三面板英文图，论文用）
"""

import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.stdout.reconfigure(encoding="utf-8")

from music21 import converter

from scoreflow.analysis.cross_hand import reassign
from scoreflow.analysis.engine import AnalysisEngine, ScoreResult
from scoreflow.analysis.features import build_features
from scoreflow.config import load_config
from human_vs_auto_compare import load_human

# 引擎 v1.4 特征名（wide_leap_sum 已更名为 stretched_grip，与 build_features 一致）
FEATURE_ORDER = [
    "stretched_grip", "contracted_grip", "consecutive_sixteenths",
    "consecutive_eighths", "consecutive_chords", "consecutive_intervals",
]


def merge_pages(results, piece: str) -> ScoreResult:
    """合并多页 ScoreResult：特征小节相加、难度分求和、难度明细聚合。"""
    merged = ScoreResult(source=Path(f"{piece}.pdf"), pages=len(results))
    for name in results[0].feature_results:
        total = sum(r.feature_results[name].total_measures for r in results)
        matched = sum(r.feature_results[name].matched_measures for r in results)
        hits = []
        for r in results:
            hits.extend(r.feature_results[name].hits)
        merged.feature_results[name] = type(results[0].feature_results[name])(
            feature=results[0].feature_results[name].feature,
            total_measures=total,
            matched_measures=matched,
            hits=hits,
        )
    merged.difficulty_score = round(sum(r.difficulty_score for r in results), 2)
    for name in results[0].feature_results:
        agg = {"weight": 0.0, "hit_measures": 0, "two_hand_measures": 0, "score": 0.0}
        for r in results:
            rd = r.difficulty_detail.get(name)
            if not rd:
                continue
            agg["weight"] = rd["weight"]
            agg["hit_measures"] += rd["hit_measures"]
            agg["two_hand_measures"] += rd["two_hand_measures"]
            agg["score"] += rd["score"]
        merged.difficulty_detail[name] = agg
    return merged

HUMAN_XLSX = Path(r"C:\Users\BinyuWang\Desktop\CV\etude.xlsx")
OUTPUT_ROOT = Path(r"D:\HOMR\output")
OUT_DIR = Path(r"D:\HOMR\output_grades\reports")
FIG_DIR = HERE / "cvaa_figs_en"

PIECES = [f"etude{i:02d}" for i in range(1, 11)]


def xml_pages(piece: str) -> list[Path]:
    return sorted((OUTPUT_ROOT / piece / "pages").glob("page_*.musicxml"))


def run_config(pages_by_piece, features, chord_mode,
               hand_split, high, low) -> dict:
    """对一个配置跑全语料 -> {piece: {"matched": {feat: n}, "difficulty": d}}"""
    engine = AnalysisEngine(
        features=features,
        chord_mode=chord_mode,
        cross_hand_heuristic=False,   # 交叉手重分配在这里手动做（便于扫阈值）
        hand_split_threshold=hand_split,
    )
    out = {}
    for piece in pages_by_piece:
        results = []
        for xp in pages_by_piece[piece]:
            text = xp.read_text(encoding="utf-8", errors="ignore")
            text, _ = reassign(text, high_threshold=high, low_threshold=low)
            score = converter.parseData(text)
            results.append(engine._maybe_split_and_analyze(score, source=xp))
        merged = merge_pages(results, piece)
        out[piece] = {
            "matched": {f: merged.feature_results[f].matched_measures
                        for f in FEATURE_ORDER},
            "difficulty": merged.difficulty_score,
        }
    return out


def agreement(auto, human) -> tuple[float, dict]:
    """60 单元一致率均值 + 分特征均值。"""
    # 人工 xlsx 的特征 key 仍用旧名 wide_leap_sum（引擎已更名 stretched_grip）
    human_key = {f: ("wide_leap_sum" if f == "stretched_grip" else f)
                 for f in FEATURE_ORDER}
    per_feat = {f: [] for f in FEATURE_ORDER}
    for g, entry in human.items():
        piece = f"etude{g:02d}"
        if piece not in auto:
            continue
        n = entry["小节数"]
        for f in FEATURE_ORDER:
            h = entry.get(human_key[f], 0)
            a = auto[piece]["matched"][f]
            per_feat[f].append(1 - abs(a - h) / n)
    means = {f: sum(v) / len(v) for f, v in per_feat.items()}
    overall = sum(means.values()) / len(means)
    return overall, means


def spearman(grades, scores) -> float:
    def rank(v):
        order = sorted(range(len(v)), key=lambda i: v[i])
        r = [0.0] * len(v)
        i = 0
        while i < len(order):
            j = i
            while j + 1 < len(order) and v[order[j + 1]] == v[order[i]]:
                j += 1
            avg = (i + j) / 2 + 1
            for k in range(i, j + 1):
                r[order[k]] = avg
            i = j + 1
        return r
    n = len(grades)
    rx, ry = rank(grades), rank(scores)
    mx, my = sum(rx) / n, sum(ry) / n
    num = sum((a - mx) * (b - my) for a, b in zip(rx, ry))
    den = (sum((a - mx) ** 2 for a in rx) * sum((b - my) ** 2 for b in ry)) ** 0.5
    return num / den if den else 0.0


def evaluate(auto, human):
    agree, per_feat = agreement(auto, human)
    grades = sorted(human)
    rho = spearman(grades, [auto[f"etude{g:02d}"]["difficulty"] for g in grades])
    return agree, per_feat, rho


def main():
    t0 = time.time()
    cfg = load_config()
    features = build_features(cfg.features)
    chord_mode = cfg.chord_mode
    human = load_human(str(HUMAN_XLSX))
    pages_by_piece = {p: xml_pages(p) for p in PIECES}
    assert all(v for v in pages_by_piece.values()), "缺少 OMR 页面 MusicXML"

    sweeps = {
        "hand_split": [12, 13, 14, 15, 16, 17, 18, 19, 20],
        "cross_high": [58, 60, 63, 65, 67, 69, 72, 74, 76, 79],
        "cross_low": [36, 39, 41, 43, 45, 48, 50, 52, 55],
    }
    results = {"hand_split": [], "cross_high": [], "cross_low": []}
    total = sum(len(v) for v in sweeps.values())
    done = 0

    for t in sweeps["hand_split"]:
        auto = run_config(pages_by_piece, features, chord_mode,
                          hand_split=t, high=72, low=48)
        agree, per, rho = evaluate(auto, human)
        results["hand_split"].append((t, agree, per, rho))
        done += 1
        print(f"[{done}/{total}] hand_split={t}  agree={agree*100:.1f}%  rho={rho:.3f}  "
              f"({time.time()-t0:.0f}s)", flush=True)

    for h in sweeps["cross_high"]:
        auto = run_config(pages_by_piece, features, chord_mode,
                          hand_split=None, high=h, low=48)
        agree, per, rho = evaluate(auto, human)
        results["cross_high"].append((h, agree, per, rho))
        done += 1
        print(f"[{done}/{total}] cross_high={h}  agree={agree*100:.1f}%  rho={rho:.3f}  "
              f"({time.time()-t0:.0f}s)", flush=True)

    for l in sweeps["cross_low"]:
        auto = run_config(pages_by_piece, features, chord_mode,
                          hand_split=None, high=72, low=l)
        agree, per, rho = evaluate(auto, human)
        results["cross_low"].append((l, agree, per, rho))
        done += 1
        print(f"[{done}/{total}] cross_low={l}  agree={agree*100:.1f}%  rho={rho:.3f}  "
              f"({time.time()-t0:.0f}s)", flush=True)

    # ---------- Markdown 报告 ----------
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    L = ["# 阈值敏感性分析（回应外审意见 2）", "",
         f"- 生成时间: {time.strftime('%Y-%m-%d %H:%M:%S')}",
         "- 语料: 中国音乐学院考级 1-10 级 10 首（36 页、400 小节）",
         "- 指标: 60 单元一致率（c = 1 - |a-h|/n）；难度分与等级的 Spearman ρ",
         "- 方法: one-at-a-time 扫描，其余参数固定为论文默认值", ""]

    titles = {
        "hand_split": ("双手拆分阈值（半音，默认 16；配置含交叉手重分配 72/48）", sweeps["hand_split"], 16),
        "cross_high": ("交叉手重分配高阈值（bass 谱表 > 阈值归右手，默认 72 = C5；论文主表口径，无双手拆分）", sweeps["cross_high"], 72),
        "cross_low": ("交叉手重分配低阈值（treble 谱表 < 阈值归左手，默认 48 = C3；论文主表口径，无双手拆分）", sweeps["cross_low"], 48),
    }
    for key, (title, values, default) in titles.items():
        L.append(f"## {title}")
        L.append("")
        L.append("| 阈值 | 默认 | 60 单元一致率 | 分特征最差 | Spearman ρ |")
        L.append("|---:|:---:|---:|---:|---:|")
        for v, agree, per, rho in results[key]:
            worst_feat = min(per, key=per.get)
            L.append(f"| {v} | {'**是**' if v == default else ''} "
                     f"| {agree*100:.1f}% | {per[worst_feat]*100:.1f}% ({worst_feat}) "
                     f"| {rho:.3f} |")
        L.append("")

    agree_span = {}
    for key in results:
        vals = [a for _, a, _, _ in results[key]]
        agree_span[key] = (min(vals), max(vals))
    L.append("## 结论摘要")
    L.append("")
    for key, (mn, mx) in agree_span.items():
        L.append(f"- {titles[key][0]}：一致率范围 {mn*100:.1f}% – {mx*100:.1f}%"
                 f"（极差 {(mx-mn)*100:.1f} 个百分点）")

    md_path = OUT_DIR / "敏感性分析.md"
    md_path.write_text("\n".join(L), encoding="utf-8")
    print(f"✓ {md_path}", flush=True)

    # ---------- Excel ----------
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor="4472C4")
    hdr_font = Font(bold=True, color="FFFFFF")
    row0 = 1
    for key, (title, values, default) in titles.items():
        ws = wb.create_sheet(key[:31])
        ws.append([title])
        ws.append(["阈值", "是否默认", "60单元一致率%", "Spearman ρ",
                   "伸张把位%", "收缩把位%", "连续十六分%", "连续八分%",
                   "连续和弦%", "连续音程%"])
        for v, agree, per, rho in results[key]:
            ws.append([v, "是" if v == default else "",
                       round(agree * 100, 1), round(rho, 3)] +
                      [round(per[f] * 100, 1) for f in FEATURE_ORDER])
        for c in range(1, 11):
            cell = ws.cell(row=2, column=c)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions["A"].width = 10
    ws_sum = wb.create_sheet("结论")
    ws_sum.append(["扫描维度", "一致率最小%", "一致率最大%", "极差(百分点)"])
    for key, (mn, mx) in agree_span.items():
        ws_sum.append([titles[key][0], round(mn * 100, 1), round(mx * 100, 1),
                       round((mx - mn) * 100, 1)])
    xlsx_path = OUT_DIR / "敏感性分析.xlsx"
    wb.save(xlsx_path)
    print(f"✓ {xlsx_path}", flush=True)

    # ---------- 英文图（三面板） ----------
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.sans-serif"] = ["DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False

    FIG_DIR.mkdir(exist_ok=True)
    panels = [
        ("hand_split", "Hand-split threshold (semitones)", sweeps["hand_split"], 16),
        ("cross_high", "Cross-hand high threshold (MIDI, default C5 = 72)", sweeps["cross_high"], 72),
        ("cross_low", "Cross-hand low threshold (MIDI, default C3 = 48)", sweeps["cross_low"], 48),
    ]
    fig, axes = plt.subplots(1, 3, figsize=(11.5, 3.4), dpi=200)
    for ax, (key, label, values, default) in zip(axes, panels):
        vs = [v for v, _, _, _ in results[key]]
        ags = [a * 100 for _, a, _, _ in results[key]]
        rhos = [r for _, _, _, r in results[key]]
        ax.plot(vs, ags, "o-", color="#2E75B6", lw=2, ms=6,
                label="60-cell agreement (%)")
        ax.axvline(default, color="#C00000", ls="--", lw=1.2,
                   label=f"paper default = {default}")
        ax.set_xlabel(label, fontsize=9)
        ax.set_ylabel("Agreement rate (%)", fontsize=9, color="#185FA5")
        ax.tick_params(axis="both", labelsize=8)
        ax2 = ax.twinx()
        ax2.plot(vs, rhos, "s--", color="#639922", lw=1.4, ms=5,
                 label="Spearman ρ")
        ax2.set_ylabel("Spearman ρ", fontsize=9, color="#3B6D11")
        ax2.set_ylim(0.90, 1.01)
        ax2.tick_params(axis="y", labelsize=8)
        ax.set_ylim(88.0, 94.5)
        ax.set_title(label.split("(")[0].strip(), fontsize=10)
        h1, l1 = ax.get_legend_handles_labels()
        h2, l2 = ax2.get_legend_handles_labels()
        ax.legend(h1 + h2, l1 + l2, fontsize=7, loc="lower left",
                  framealpha=0.9)
        ax.grid(alpha=0.3, ls=":")
    fig.tight_layout()
    fig_path = FIG_DIR / "fig_sensitivity_en.png"
    fig.savefig(fig_path, bbox_inches="tight", facecolor="white")
    print(f"✓ {fig_path}", flush=True)
    print(f"DONE in {time.time()-t0:.0f}s", flush=True)


if __name__ == "__main__":
    main()
